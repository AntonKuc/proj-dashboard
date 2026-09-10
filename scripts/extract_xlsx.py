import openpyxl, json, datetime, sys, os

# Put the source report here when you want to regenerate data/raw_extract.json
# (not committed to git - it's your raw financial data, see .gitignore).
SRC = os.environ.get(
    'SOURCE_XLSX',
    os.path.join(os.path.dirname(__file__), '..', 'data', 'source', 'report.xlsx'),
)
OUT = os.path.join(os.path.dirname(__file__), '..', 'data', 'raw_extract.json')

skip_substrings = ['коэфф', 'в среднем']

# Months from Sep 2026 onward are placeholder zeros in the source sheet
# (the report was last actually filled through August 2026) - exclude them
# so we don't store fabricated "zero revenue" facts for unreported months.
LAST_REPORTED_MONTH = '2026-08'

def main():
    wb = openpyxl.load_workbook(SRC, data_only=True)
    records = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = list(ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True))
        if not rows:
            continue
        # header row is the first row containing 'Наименование' in column B
        header_idx = None
        for i, row in enumerate(rows):
            if len(row) > 1 and isinstance(row[1], str) and 'наименование' in row[1].lower():
                header_idx = i
                break
        if header_idx is None:
            print(f'WARNING: no header row found in sheet {sheet_name}', file=sys.stderr)
            continue
        header = rows[header_idx]
        month_cols = []
        for idx, cell in enumerate(header):
            if isinstance(cell, datetime.datetime):
                month_cols.append((idx, cell.strftime('%Y-%m')))
        for row in rows[header_idx + 1:]:
            if row is None:
                continue
            name = row[1] if len(row) > 1 else None
            if not name or not isinstance(name, str):
                continue
            lname = name.lower()
            if any(s in lname for s in skip_substrings):
                continue
            for idx, month in month_cols:
                if month > LAST_REPORTED_MONTH:
                    continue
                if idx >= len(row):
                    continue
                val = row[idx]
                if val is None:
                    continue
                if not isinstance(val, (int, float)):
                    continue
                records.append({
                    'projectCode': sheet_name,
                    'line': name.strip(),
                    'month': month,
                    'value': float(val)
                })
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(records, f, ensure_ascii=False, indent=2)
    print(f'Wrote {len(records)} records to {OUT}')

    # sanity: unique project codes and lines
    projects = sorted(set(r['projectCode'] for r in records))
    lines = sorted(set(r['line'] for r in records))
    print('Projects:', projects)
    print('Distinct line count:', len(lines))

if __name__ == '__main__':
    main()
